"""Best-effort import of an uploaded DOCX/PDF resume into the master structure (FR-002).

Nothing here is authoritative: the user reviews every field before saving.
"""

import io
import re
from dataclasses import dataclass

from jobhunter.services.resume.model import empty

HEADINGS = {
    "summary": (
        "summary",
        "profile",
        "professional summary",
        "about me",
        "career summary",
        "objective",
        "professional profile",
    ),
    "experience": (
        "experience",
        "work experience",
        "professional experience",
        "employment",
        "employment history",
        "work history",
        "career history",
    ),
    "education": ("education", "education and training", "academic background"),
    "certifications": (
        "certifications",
        "certification",
        "certificates",
        "licenses",
        "licenses & certifications",
        "licenses and certifications",
        "certifications and training",
        "professional development",
    ),
    "skills": (
        "skills",
        "technical skills",
        "core competencies",
        "key skills",
        "competencies",
        "skills & tools",
        "tools",
        "technologies",
        "areas of expertise",
    ),
}
_HEADING_LOOKUP = {v: k for k, vs in HEADINGS.items() for v in vs}
_MONTH = r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?"
_DATE = rf"(?:{_MONTH}\s+\d{{4}}|\d{{1,2}}/\d{{4}}|\d{{4}})"
DATE_RANGE = re.compile(
    rf"(?P<start>{_DATE})\s*(?:–|—|-|to)\s*(?P<end>{_DATE}|present|current|now|today)",
    re.IGNORECASE,
)
_EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
_PHONE = re.compile(r"(?:\+?1[\s.-]?)?\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}")
_LINK = re.compile(r"(?:https?://|www\.|linkedin\.com/|github\.com/)\S+", re.IGNORECASE)
_CITY = re.compile(r"\b([A-Z][a-zA-Z.'-]+(?:\s[A-Z][a-zA-Z.'-]+)*),\s*([A-Z]{2})\b")
_BULLET = re.compile(r"^\s*[•●▪◦\-*·]\s*")
_SPLITS = (" at ", " | ", " — ", " – ", " - ", ", ")
_DEGREE = re.compile(
    r"\b(bachelor|master|b\.?sc|m\.?sc|b\.?a|m\.?a|ph\.?d|diploma|certificate|"
    r"degree|b\.?eng|m\.?eng|mba|associate)\b",
    re.IGNORECASE,
)
_SCHOOL = re.compile(
    r"\b(university|college|institute|school|polytechnic|academy|sait|nait|"
    r"bcit)\b",
    re.IGNORECASE,
)


@dataclass
class Line:
    text: str
    heading: bool = False
    bullet: bool = False


class ImportError_(ValueError):
    pass


def lines_from_docx(data: bytes) -> list[Line]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    out = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        is_list = "list" in style or p._p.pPr is not None and p._p.pPr.numPr is not None
        out.append(
            Line(
                _BULLET.sub("", text),
                heading=style.startswith(("heading", "title")),
                bullet=is_list or bool(_BULLET.match(text)),
            )
        )
    for table in doc.tables:  # some resumes use tables for layout
        for row in table.rows:
            for cell in row.cells:
                for para in cell.text.splitlines():
                    if para.strip():
                        out.append(
                            Line(_BULLET.sub("", para.strip()), bullet=bool(_BULLET.match(para)))
                        )
    return out


def lines_from_pdf(data: bytes) -> list[Line]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    return [
        Line(_BULLET.sub("", t.strip()), bullet=bool(_BULLET.match(t)))
        for t in text.splitlines()
        if t.strip()
    ]


def _section_of(line: Line) -> str | None:
    key = re.sub(r"[^a-z& ]", "", line.text.lower()).strip()
    if len(key) > 40:
        return None
    return _HEADING_LOOKUP.get(key)


def _title_employer(text: str) -> tuple[str, str]:
    for sep in _SPLITS:
        if sep in text:
            left, right = text.split(sep, 1)
            return left.strip(" ,|–—-"), right.strip(" ,|–—-")
    return text.strip(), ""


def _parse_experience(lines: list[Line]) -> list[dict]:
    entries: list[dict] = []
    current: dict | None = None
    for i, line in enumerate(lines):
        dates = DATE_RANGE.search(line.text)
        nxt = lines[i + 1] if i + 1 < len(lines) else None
        if line.bullet and current is not None:
            current["bullets"].append(line.text)
            continue
        if dates:
            rest = (line.text[: dates.start()] + line.text[dates.end() :]).strip(" ,|–—-()")
            if current is not None and not current["start"] and not current["bullets"]:
                current["start"], current["end"] = dates.group("start"), dates.group("end")
                if rest and not current["location"]:
                    current["location"] = rest
                continue
            title, employer = _title_employer(rest) if rest else ("", "")
            current = {
                "employer": employer,
                "title": title,
                "location": "",
                "start": dates.group("start"),
                "end": dates.group("end"),
                "bullets": [],
            }
            entries.append(current)
            continue
        starts_entry = (
            current is None
            or current["bullets"]
            or (nxt is not None and DATE_RANGE.search(nxt.text) and not nxt.bullet)
        )
        if starts_entry and not line.bullet:
            title, employer = _title_employer(line.text)
            current = {
                "employer": employer,
                "title": title,
                "location": "",
                "start": "",
                "end": "",
                "bullets": [],
            }
            entries.append(current)
        elif current is not None and not current["employer"] and not current["bullets"]:
            current["employer"] = line.text
        elif (
            current is not None
            and not current["location"]
            and not current["bullets"]
            and _CITY.search(line.text)
        ):
            current["location"] = line.text
        elif current is not None:
            current["bullets"].append(line.text)
    return entries


def _parse_education(lines: list[Line]) -> list[dict]:
    entries: list[dict] = []
    current: dict | None = None
    for line in lines:
        dates = DATE_RANGE.search(line.text) or re.search(r"\b(19|20)\d{2}\b", line.text)
        text = re.sub(DATE_RANGE, "", line.text).strip(" ,|–—-")
        text = re.sub(r"\b(19|20)\d{2}\b", "", text).strip(" ,|–—-")
        school = _SCHOOL.search(text)
        degree = _DEGREE.search(text)
        if (
            current is None
            or (school and current["institution"])
            or (degree and current["credential"])
        ):
            current = {"institution": "", "credential": "", "field": "", "start": "", "end": ""}
            entries.append(current)
        if school and degree:
            cred, inst = _title_employer(text)
            if _SCHOOL.search(cred):
                cred, inst = inst, cred
            current["credential"], current["institution"] = cred, inst
        elif school:
            current["institution"] = text
        else:
            current["credential"] = current["credential"] or text
        if dates is not None:
            if isinstance(dates, re.Match) and dates.re is DATE_RANGE:
                current["start"], current["end"] = dates.group("start"), dates.group("end")
            else:
                current["end"] = dates.group(0)
    return [e for e in entries if e["institution"] or e["credential"]]


def parse_lines(lines: list[Line]) -> dict:
    data = empty()
    sections: dict[str, list[Line]] = {}
    header: list[Line] = []
    current: str | None = None
    for line in lines:
        section = _section_of(line)
        if section:
            current = section
            sections.setdefault(section, [])
            continue
        if current is None:
            header.append(line)
        else:
            sections[current].append(line)

    header_text = "\n".join(line.text for line in header)
    if header:
        data["name"] = header[0].text if not _EMAIL.search(header[0].text) else ""
    if m := _EMAIL.search(header_text):
        data["email"] = m.group(0)
    if m := _PHONE.search(header_text):
        data["phone"] = m.group(0)
    if m := _CITY.search("\n".join(line.text for line in header[1:])):
        data["location"] = m.group(0)
    data["links"] = list(dict.fromkeys(_LINK.findall(header_text)))

    data["summary"] = " ".join(line.text for line in sections.get("summary", []))
    data["experience"] = _parse_experience(sections.get("experience", []))
    data["education"] = _parse_education(sections.get("education", []))
    data["certifications"] = [
        {"name": re.sub(DATE_RANGE, "", line.text).strip(" ,|–—-"), "issuer": "", "date": ""}
        for line in sections.get("certifications", [])
    ]
    skills: list[str] = []
    for line in sections.get("skills", []):
        for part in re.split(r"[,;|•]", line.text):
            part = re.sub(r"^[^:]{1,25}:\s*", "", part).strip()  # "Tools: Jira" -> "Jira"
            if part:
                skills.append(part)
    data["skills"] = list(dict.fromkeys(skills))
    return data


def import_resume(data: bytes, fmt: str) -> dict:
    lines = lines_from_docx(data) if fmt == "docx" else lines_from_pdf(data)
    if not lines:
        raise ImportError_("No text could be read from this file (is it a scanned image?).")
    return parse_lines(lines)
