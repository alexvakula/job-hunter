"""ATS-safe resume and cover letter as DOCX and PDF (FR-009, FR-010).

One column, standard headings, plain paragraphs only: no tables, images, text boxes or
header/footer text. The look follows the family's reference resume: name, contact lines and a
headline on top; navy upper-case section headings with a rule; "Employer, Location  dates" over
an italic job title; "- " bullets. PDFs use bundled DejaVu fonts so any Unicode text renders.
"""

import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from fpdf import FPDF

FONTS = Path(__file__).resolve().parents[2] / "fonts"
HEADINGS = ("SUMMARY", "SKILLS", "WORK EXPERIENCE", "EDUCATION", "CERTIFICATIONS")
NAVY, GREY, LINK, BLACK = "1A3E6D", "4A5568", "0563C1", "000000"
FONT = "Calibri"


def _dates(e: dict) -> str:
    start, end = e.get("start", ""), e.get("end", "")
    return f"{start} – {end}" if start and end else start or end


def _url(link: str) -> str:
    return link if "://" in link else f"https://{link}"


def _contact_lines(c: dict) -> list[list[tuple[str, str]]]:
    """Two lines of (text, url) pieces: "Location, Tel.: phone" and "email | link | link"."""
    phone = f"Tel.: {c['phone']}" if c.get("phone") else ""
    first = ", ".join(p for p in (c.get("location"), phone) if p)
    second: list[tuple[str, str]] = []
    for text, url in [(c.get("email"), "")] + [(x, _url(x)) for x in c.get("links", [])]:
        if text:
            second += [(" | ", "")] if second else []
            second.append((text, url))
    return [line for line in ([(first, "")] if first else [], second) if line]


def _skill_groups(c: dict) -> list[dict]:
    """Snapshots from before skill groups existed only have the flat list."""
    if "skill_groups" in c:
        return c["skill_groups"]
    return [{"label": "", "skills": c["skills"]}] if c.get("skills") else []


def _blocks(c: dict) -> list[tuple[str, object]]:
    """The resume body as (kind, value) blocks, shared by the DOCX and PDF writers.

    Kinds: heading, text, labeled (label, text), entry (employer, location, dates), role, bullet.
    """
    out: list[tuple[str, object]] = []
    if c.get("summary"):
        out += [("heading", "SUMMARY"), ("text", c["summary"])]
    groups = _skill_groups(c)
    if groups:
        out.append(("heading", "SKILLS"))
        out += [("labeled", (g["label"], ", ".join(g["skills"]))) for g in groups]
    if c.get("experience"):
        out.append(("heading", "WORK EXPERIENCE"))
        for e in c["experience"]:
            out.append(("entry", (e.get("employer", ""), e.get("location", ""), _dates(e))))
            if e.get("title"):
                out.append(("role", e["title"]))
            out += [("bullet", b) for b in e.get("bullets", [])]
    if c.get("education"):
        out.append(("heading", "EDUCATION"))
        for e in c["education"]:
            school = ", ".join(p for p in (e.get("institution"), _dates(e)) if p)
            cred, field = e.get("credential", ""), e.get("field", "")
            if field and field.lower() not in cred.lower():
                cred = ", ".join(p for p in (cred, field) if p)
            out += [("text", t) for t in (school, cred) if t]
    if c.get("certifications"):
        out.append(("heading", "CERTIFICATIONS"))
        for x in c["certifications"]:
            text = x.get("name", "") + (f" - {x['issuer']}" if x.get("issuer") else "")
            out.append(("bullet", ", ".join(p for p in (text, x.get("date")) if p)))
    return out


# --- DOCX ----------------------------------------------------------------------------------

# w:pPr children that must follow w:pBdr (ECMA-376 order), so Word accepts the border.
_AFTER_PBDR = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
    "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
    "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
    "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
    "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)  # fmt: skip


def _new_doc() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.75)
    sec.top_margin = sec.bottom_margin = Inches(0.625)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(10.5)
    normal.paragraph_format.space_after = Pt(0)
    h = doc.styles["Heading 1"]
    h.font.name, h.font.size, h.font.color.rgb = FONT, Pt(11), RGBColor.from_string(NAVY)
    h.font.bold, h.font.italic, h.font.underline = True, False, True
    rfonts = h.element.rPr.rFonts
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(13), Pt(5)
    h.paragraph_format.keep_with_next = True
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in {"w:val": "single", "w:sz": "4", "w:space": "2", "w:color": NAVY}.items():
        bottom.set(qn(k), v)
    border.append(bottom)
    h.element.get_or_add_pPr().insert_element_before(border, *_AFTER_PBDR)
    return doc


def _para(doc, before=0.0, after=0.0, line=None, indent=0.0, hanging=0.0, keep=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if line:
        pf.line_spacing_rule, pf.line_spacing = WD_LINE_SPACING.MULTIPLE, line
    if indent:
        pf.left_indent, pf.first_line_indent = Inches(indent + hanging), Inches(-hanging)
    pf.keep_with_next = keep or None
    return p


def _run(p, text, size=10.5, bold=False, italic=False, color=None):
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold, r.italic = bold or None, italic or None
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def _link(p, text, url, size):
    r = _run(p, text, size, color=LINK)
    r.font.underline = True
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    h.append(r._r)  # moves the run inside the hyperlink
    p._p.append(h)


def _docx_header(doc, c: dict, headline: bool):
    _run(_para(doc, after=2), c.get("name") or "", 16, bold=True)
    lines = _contact_lines(c)
    for i, pieces in enumerate(lines):
        p = _para(doc, after=5 if i == len(lines) - 1 else 2)
        for text, url in pieces:
            _link(p, text, url, 10) if url else _run(p, text, 10)
    if headline and c.get("headline"):
        _run(_para(doc, after=11), c["headline"], 12, bold=True, color=NAVY)


def resume_docx(c: dict) -> bytes:
    doc = _new_doc()
    _docx_header(doc, c, headline=True)
    for kind, value in _blocks(c):
        if kind == "heading":
            doc.add_paragraph(value, style="Heading 1")
        elif kind == "text":
            _run(_para(doc, after=5.5, line=1.1), value)
        elif kind == "labeled":
            label, text = value
            p = _para(doc, after=3.5, line=1.12)
            if label:
                _run(p, f"{label}: ", bold=True)
            _run(p, text)
        elif kind == "entry":
            employer, location, dates = value
            p = _para(doc, before=9, after=1, keep=True)
            _run(p, employer, 11, bold=True)
            if location:
                _run(p, f", {location}", 11)
            if dates:
                _run(p, f"   {dates}", italic=True, color=GREY)
        elif kind == "role":
            _run(_para(doc, after=5, keep=True), value, italic=True)
        else:
            _run(_para(doc, after=2.5, line=1.06, indent=0.25, hanging=0.12), f"- {value}")
    return _save(doc)


def letter_docx(c: dict, letter: str) -> bytes:
    doc = _new_doc()
    _docx_header(doc, c, headline=False)
    _para(doc, after=6)
    for para in _paragraphs(letter):
        _run(_para(doc, after=8, line=1.1), para, 11)
    return _save(doc)


def _save(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _paragraphs(text: str) -> list[str]:
    paras, current = [], []
    for line in (text or "").splitlines():
        if line.strip():
            current.append(line.strip())
        elif current:
            paras.append(" ".join(current))
            current = []
    if current:
        paras.append(" ".join(current))
    return paras


# --- PDF -----------------------------------------------------------------------------------
# DejaVu Sans runs wider than Calibri, so PDF sizes are about one point smaller than the DOCX.


def _rgb(hex_: str) -> tuple[int, int, int]:
    return int(hex_[:2], 16), int(hex_[2:4], 16), int(hex_[4:], 16)


class _Pdf(FPDF):
    def __init__(self):
        super().__init__(format="Letter", unit="pt")
        self.set_margins(54, 45, 54)
        self.set_auto_page_break(True, margin=45)
        self.add_font("DejaVu", "", str(FONTS / "DejaVuSans.ttf"))
        self.add_font("DejaVu", "B", str(FONTS / "DejaVuSans-Bold.ttf"))
        self.add_font("DejaVu", "I", str(FONTS / "DejaVuSans-Oblique.ttf"))
        self.add_page()

    def _style(self, size: float, style: str = "", color: str = BLACK):
        self.set_font("DejaVu", style, size)
        self.set_text_color(*_rgb(color))

    def keep_room(self, height: float):
        """Start a new page rather than leave a heading or job title alone at the bottom."""
        if self.get_y() + height > self.page_break_trigger:
            self.add_page()

    def pieces(self, parts, size: float, before=0.0, after=0.0):
        """One paragraph of mixed (text, style, size, color, url) runs; wraps at the margin."""
        self.ln(before)
        for text, style, run_size, color, url in parts:
            self._style(run_size, style, color)
            self.write(size * 1.2, text, link=url)
        self.ln(size * 1.2 + after)

    def block(self, text: str, size: float, style="", color=BLACK, after=0.0, lead=1.25):
        self._style(size, style, color)
        self.multi_cell(0, size * lead, text, align="L", new_x="LMARGIN", new_y="NEXT")
        self.ln(after)

    def bullet(self, text: str, size: float, indent=18.0, hang=9.0, after=1.5):
        self._style(size)
        self.set_x(self.l_margin + indent)
        self.cell(hang, size * 1.25, "-")
        self.multi_cell(0, size * 1.25, text, align="L", new_x="LMARGIN", new_y="NEXT")
        self.ln(after)

    def heading(self, text: str):
        self.ln(11)
        self.keep_room(60)
        self.block(text, 10.5, "BU", NAVY)
        y = self.get_y() + 1.5
        self.set_draw_color(*_rgb(NAVY))
        self.set_line_width(0.5)
        self.line(self.l_margin, y, self.w - self.r_margin, y)
        self.ln(6)

    def header_block(self, c: dict, headline: bool):
        self.block(c.get("name") or "", 15, "B", after=3)
        for pieces in _contact_lines(c):
            self.pieces(
                [(t, "U" if u else "", 9, LINK if u else BLACK, u) for t, u in pieces], 9, after=1
            )
        if headline and c.get("headline"):
            self.ln(3)
            self.block(c["headline"], 11, "B", NAVY, after=6)


def resume_pdf(c: dict) -> bytes:
    pdf = _Pdf()
    pdf.header_block(c, headline=True)
    for kind, value in _blocks(c):
        if kind == "heading":
            pdf.heading(value)
        elif kind == "text":
            pdf.block(value, 9.5, after=4)
        elif kind == "labeled":
            label, text = value
            bold = [(f"{label}: ", "B", 9.5, BLACK, "")] if label else []
            pdf.pieces(bold + [(text, "", 9.5, BLACK, "")], 9.5, after=2)
        elif kind == "entry":
            employer, location, dates = value
            pdf.keep_room(50)
            parts = [(employer, "B", 10, BLACK, "")]
            parts += [(f", {location}", "", 10, BLACK, "")] if location else []
            parts += [(f"   {dates}", "I", 9.5, GREY, "")] if dates else []
            pdf.pieces(parts, 10, before=7)
        elif kind == "role":
            pdf.block(value, 9.5, "I", after=4)
        else:
            pdf.bullet(value, 9.5)
    return bytes(pdf.output())


def letter_pdf(c: dict, letter: str) -> bytes:
    pdf = _Pdf()
    pdf.header_block(c, headline=False)
    pdf.ln(12)
    for para in _paragraphs(letter):
        pdf.block(para, 10.5, after=7)
    return bytes(pdf.output())
