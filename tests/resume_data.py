"""Sample master resume, posting and DOCX builder for feature 004 tests."""

import io

POSTING = """QA Lead — Acme Robotics (Calgary, AB, hybrid)

We are looking for a QA Lead to own our test strategy and grow test automation across web and
mobile. You will mentor a team of 4 testers and work closely with stakeholders.

Requirements:
- 5+ years in quality assurance, 2+ years leading a team
- Strong test automation with Selenium or Playwright, and API testing with Postman
- Experience with Jira, CI/CD (Jenkins or GitHub Actions) and Agile/Scrum
- Performance testing with JMeter is an asset
- ISTQB certification preferred
"""

MASTER = {
    "name": "Sam Rivera",
    "email": "sam.jobs@example.org",
    "phone": "403-555-0100",
    "location": "Calgary, AB",
    "links": ["https://linkedin.com/in/samrivera"],
    "summary": "QA leader with 10 years of experience building test automation and leading teams.",
    "experience": [
        {
            "id": "e1",
            "employer": "Northwind Imaging",
            "title": "QA Lead",
            "location": "Calgary, AB",
            "start": "Jan 2020",
            "end": "Present",
            "bullets": [
                "Introduced quarterly release reviews with product owners.",
                "Built a Selenium and Python test automation framework covering 1,200 cases.",
                "Ran API testing with Postman and integrated suites into Jenkins CI/CD.",
                "Mentored 5 testers; tracked defects in Jira.",
            ],
        },
        {
            "id": "e2",
            "employer": "Société Générale Tech",
            "title": "Senior QA Analyst",
            "location": "Montréal, QC",
            "start": "2015",
            "end": "2019",
            "bullets": [
                "Wrote test plans and test cases for trading systems.",
                "Worked in Agile Scrum teams of 8.",
            ],
        },
    ],
    "education": [
        {
            "institution": "University of Calgary",
            "credential": "BSc",
            "field": "Computer Science",
            "start": "2008",
            "end": "2012",
        }
    ],
    "certifications": [{"name": "ISTQB Foundation", "issuer": "ISTQB", "date": "2014"}],
    "skills": ["Leadership", "Selenium", "Python", "Jira", "Postman", "Jenkins", "Agile", "SQL"],
}


def resume_docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Sam Rivera")
    doc.add_paragraph(
        "sam.jobs@example.org | 403-555-0100 | Calgary, AB | linkedin.com/in/samrivera"
    )
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph("QA leader with 10 years of experience.")
    doc.add_heading("Work Experience", level=1)
    doc.add_paragraph("QA Lead — Northwind Imaging")
    doc.add_paragraph("Calgary, AB | Jan 2020 – Present")
    doc.add_paragraph("Built a Selenium test automation framework.", style="List Bullet")
    doc.add_paragraph("Mentored 5 testers.", style="List Bullet")
    doc.add_paragraph("Senior QA Analyst, Société Générale Tech 2015 - 2019")
    doc.add_paragraph("• Wrote test plans for trading systems.")
    doc.add_heading("Education", level=1)
    doc.add_paragraph("BSc Computer Science, University of Calgary, 2012")
    doc.add_heading("Certifications", level=1)
    doc.add_paragraph("ISTQB Foundation Level")
    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph("Testing: Selenium, Postman, JMeter; Tools: Jira, Jenkins")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
