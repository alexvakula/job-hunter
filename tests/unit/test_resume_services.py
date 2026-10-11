import io

import pytest
from docx import Document
from pypdf import PdfReader

from jobhunter.services.resume import ats, importer, model, render, tailor
from tests.resume_data import MASTER, POSTING, resume_docx_bytes

# --- import --------------------------------------------------------------------------------


def test_import_docx():
    d = importer.import_resume(resume_docx_bytes(), "docx")
    assert d["name"] == "Sam Rivera"
    assert (d["email"], d["phone"], d["location"]) == (
        "sam.jobs@example.org",
        "403-555-0100",
        "Calgary, AB",
    )
    assert "linkedin.com/in/samrivera" in d["links"][0]
    assert d["summary"] == "QA leader with 10 years of experience."
    first, second = d["experience"]
    assert (first["title"], first["employer"], first["start"], first["end"]) == (
        "QA Lead",
        "Northwind Imaging",
        "Jan 2020",
        "Present",
    )
    assert first["location"] == "Calgary, AB"
    assert first["bullets"] == [
        "Built a Selenium test automation framework.",
        "Mentored 5 testers.",
    ]
    assert (second["title"], second["employer"], second["start"], second["end"]) == (
        "Senior QA Analyst",
        "Société Générale Tech",
        "2015",
        "2019",
    )
    assert second["bullets"] == ["Wrote test plans for trading systems."]
    edu = d["education"][0]
    assert edu["institution"] == "University of Calgary" and "BSc" in edu["credential"]
    assert d["certifications"][0]["name"] == "ISTQB Foundation Level"
    assert {"Selenium", "Postman", "JMeter", "Jira", "Jenkins"} <= set(d["skills"])


def test_import_pdf_roundtrip():
    pdf = render.resume_pdf(tailor.snapshot(MASTER, tailor.build_draft(MASTER, [])))
    d = importer.import_resume(pdf, "pdf")
    assert d["name"] == "Sam Rivera" and d["email"] == "sam.jobs@example.org"
    titles = [(e["title"], e["employer"]) for e in d["experience"]]
    assert ("QA Lead", "Northwind Imaging") in titles
    assert "Selenium" in d["skills"]


def test_import_empty_pdf_is_explained():
    from fpdf import FPDF

    blank = FPDF()
    blank.add_page()
    with pytest.raises(importer.ImportError_):
        importer.import_resume(bytes(blank.output()), "pdf")


def test_model_form_validation():
    form = {
        "name": "A",
        "email": "bad",
        "experience-0-employer": "X",
        "experience-0-title": "",
        "experience-1-employer": "",
        "experience-1-title": "",
        "education-0-institution": "",
        "education-0-credential": "BSc",
        "skills": "Python, SQL\nJira",
    }

    class F(dict):
        def getlist(self, k):
            return []

    data, errors = model.from_form(F(form))
    assert "email" in errors and "experience-0" in errors and "education-0" in errors
    assert data["skills"] == ["Python", "SQL", "Jira"]
    assert len(data["experience"]) == 1  # blank row dropped


# --- ATS -----------------------------------------------------------------------------------


def test_keywords_from_posting():
    kw = ats.extract_keywords(POSTING)
    for term in (
        "test strategy",
        "test automation",
        "Selenium",
        "Playwright",
        "API testing",
        "Postman",
        "Jira",
        "CI/CD",
        "Jenkins",
        "GitHub Actions",
        "Agile",
        "Scrum",
        "performance testing",
        "JMeter",
        "ISTQB",
        "mentoring",
        "team leadership",
        "stakeholder management",
    ):
        assert term in kw, term
    assert "QA" not in kw  # generic, ignored


def test_coverage():
    kw = ats.extract_keywords(POSTING)
    cov = ats.coverage(kw, model.full_text(MASTER))
    assert "Selenium" in cov.matched and "Jira" in cov.matched and "ISTQB" in cov.matched
    assert "Playwright" in cov.missing and "JMeter" in cov.missing
    assert cov.percent == round(100 * len(cov.matched) / len(kw))


def test_special_terms_boundaries():
    assert ats.present("C#", "Strong C# skills") and not ats.present("C#", "C language")
    assert ats.present("CI/CD", "continuous integration") and ats.present("Go", "Golang dev")
    assert not ats.present("Go", "go to market")
    assert not ats.present("Java", "JavaScript only")


# --- tailoring and honesty -----------------------------------------------------------------


def test_draft_orders_by_relevance_and_keeps_everything():
    kw = ats.extract_keywords(POSTING)
    draft = tailor.build_draft(MASTER, kw)
    bullets = [b["text"] for b in draft["experience"][0]["bullets"]]
    assert bullets[-1] == "Introduced quarterly release reviews with product owners."
    assert sorted(bullets) == sorted(MASTER["experience"][0]["bullets"])
    skills = [s["name"] for s in draft["skills"] if set(s) == {"name", "hidden"}]
    assert skills.index("Selenium") < skills.index("Leadership")
    assert sorted(skills) == sorted(MASTER["skills"])


def test_missing_posting_keywords_are_added_for_matching():
    kw = ats.extract_keywords(POSTING)
    draft = tailor.build_draft(MASTER, kw)
    added = [s["name"] for s in draft["skills"] if s.get("from_posting")]
    # Practices, tools and languages; never certifications (ISTQB) or leadership claims.
    assert added == [
        "test strategy",
        "quality assurance",
        "Playwright",
        "GitHub",
        "GitHub Actions",
        "performance testing",
        "JMeter",
    ]
    snap = tailor.snapshot(MASTER, draft)
    assert snap["skill_groups"] == [{"label": "", "skills": snap["skills"]}]  # one main line
    assert snap["skills"][-len(added) :] == added
    cov = ats.coverage(kw, tailor.visible_text(snap))
    assert cov.percent > ats.coverage(kw, model.full_text(MASTER)).percent
    assert cov.missing == ["stakeholder management", "team leadership"]
    # Added skills are listed, not claimed in the text, so no honesty flags.
    assert tailor.honesty_flags(MASTER, tailor.editable_text(draft, ""), kw) == []

    # Unticking one leaves it out (and the flag survives the editor form).
    i = next(i for i, s in enumerate(draft["skills"]) if s["name"] == "JMeter")
    edited = tailor.apply_form(draft, {f"s-{i}-offer": "1"})
    assert edited["skills"][i] == {"name": "JMeter", "hidden": True, "from_posting": True}
    assert "JMeter" not in tailor.snapshot(MASTER, edited)["skills"]
    assert tailor.top_up_from_posting(MASTER, edited["skills"], kw) == edited["skills"]
    ticked = tailor.apply_form(edited, {f"s-{i}-offer": "1", f"s-{i}-add": "1"})
    assert ticked["skills"][i]["hidden"] is False


def test_skill_analysis_rows_and_removal_suggestions():
    kw = [*ats.extract_keywords(POSTING), "SQL"]
    master = MASTER | {"skills": [*MASTER["skills"], "MS Office", "Python"]}
    draft = tailor.build_draft(master, kw)
    names = [s["name"] for s in draft["skills"]]
    draft["skills"][names.index("SQL")]["hidden"] = True
    snap = tailor.snapshot(master, draft)
    a = tailor.skill_analysis(master, draft["skills"], kw, tailor.visible_text(snap))
    rows = {r["term"]: r for r in a["rows"]}
    assert [r["term"] for r in a["rows"]] == kw
    assert rows["Selenium"]["state"] == "have"
    assert rows["SQL"]["state"] == "hidden"  # in the master, hidden on this resume
    assert rows["JMeter"] == {
        "term": "JMeter",
        "state": "offer",
        "index": names.index("JMeter"),
        "checked": True,
    }
    assert rows["ISTQB"]["state"] == "have"  # the sample master holds it
    assert rows["team leadership"]["state"] == "no"
    assert "leadership" in rows["team leadership"]["reason"]
    # Only skills that are neither in the posting nor QA/dev terms are suggested.
    assert [names[i] for i in a["remove"]] == ["Leadership", "MS Office"]


def test_alternative_testing_names_are_recognised():
    text = (
        "Own the regression suite, triage defects (defect triage) and keep a traceability matrix."
    )
    assert ats.extract_keywords(text) == [
        "regression testing",
        "defect management",
        "requirements traceability",
    ]
    assert ats.present("test design", "Experienced in test case design")


def test_skills_move_between_the_skills_and_exposure_boxes():
    kw = ats.extract_keywords(POSTING)
    draft = tailor.build_draft(MASTER, kw)
    names = [s["name"] for s in draft["skills"]]
    j, p = names.index("JMeter"), names.index("Python")
    form = {f"s-{j}-group": "exposure", f"s-{p}-group": "exposure", f"s-{j}-offer": "1"}
    form |= {f"s-{j}-add": "1", f"s-{p}-pos": "99"}
    moved = tailor.apply_form(draft, form)
    assert moved["skills"][-1] == {"name": "Python", "hidden": False, "exposure": True}
    snap = tailor.snapshot(MASTER, moved)
    assert snap["skill_groups"][-1] == {"label": "Exposure", "skills": ["JMeter", "Python"]}
    assert "Python" not in snap["skill_groups"][0]["skills"]
    assert snap["skills"][-2:] == ["JMeter", "Python"]  # still counted for the ATS match
    # A form without box fields keeps each skill where it was; "main" moves it back.
    assert tailor.apply_form(moved, {})["skills"][-1]["exposure"] is True
    back = tailor.apply_form(moved, {f"s-{len(names) - 1}-group": "main"})
    assert "exposure" not in back["skills"][-1]
    grouped = MASTER | {"skill_groups": [{"label": "Exposure", "skills": ["Leadership"]}]}
    snap = tailor.snapshot(grouped, moved)
    assert [g["label"] for g in snap["skill_groups"]] == ["Exposure", "Other"]
    assert snap["skill_groups"][0]["skills"] == ["Leadership", "JMeter", "Python"]


def test_posting_terms_top_up_old_drafts():
    kw = ats.extract_keywords(POSTING)
    plain = [s for s in tailor.build_draft(MASTER, kw)["skills"] if set(s) == {"name", "hidden"}]
    topped = tailor.top_up_from_posting(MASTER, plain, kw)
    assert topped[: len(plain)] == plain
    assert all(s["from_posting"] for s in topped[len(plain) :])
    grouped = MASTER | {"skill_groups": [{"label": "Tools", "skills": ["Selenium", "Jira"]}]}
    snap = tailor.snapshot(grouped, tailor.build_draft(grouped, kw))
    assert [g["label"] for g in snap["skill_groups"]] == ["Tools", "Other"]
    assert "JMeter" in snap["skill_groups"][1]["skills"]  # outside the master's groups


# --- rendering (constitution VII: DOCX structure) ------------------------------------------


def _snap():
    return tailor.snapshot(MASTER, tailor.build_draft(MASTER, ats.extract_keywords(POSTING)))


def test_docx_is_ats_safe():
    doc = Document(io.BytesIO(render.resume_docx(_snap())))
    assert (
        doc.tables == [] and doc.inline_shapes._inline_lst == []
        if hasattr(doc.inline_shapes, "_inline_lst")
        else len(doc.inline_shapes) == 0
    )
    assert len(doc.sections) == 1
    cols = doc.sections[0]._sectPr.xpath("./w:cols")
    assert not cols or cols[0].get(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num"
    ) in (None, "1")
    for section in doc.sections:
        assert all(not p.text.strip() for p in section.header.paragraphs)
        assert all(not p.text.strip() for p in section.footer.paragraphs)
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert headings == ["SUMMARY", "SKILLS", "WORK EXPERIENCE", "EDUCATION", "CERTIFICATIONS"]
    paras = [p.text for p in doc.paragraphs]
    assert "Northwind Imaging, Calgary, AB   Jan 2020 – Present" in paras
    assert (
        paras[paras.index("Northwind Imaging, Calgary, AB   Jan 2020 – Present") + 1] == "QA Lead"
    )
    assert "Société Générale Tech, Montréal, QC   2015 – 2019" in paras
    assert "- ISTQB Foundation - ISTQB, 2014" in paras


def test_pdf_text_is_extractable_and_unicode():
    reader = PdfReader(io.BytesIO(render.resume_pdf(_snap())))
    text = "\n".join(p.extract_text() for p in reader.pages)
    for needle in ("Sam Rivera", "Société Générale Tech", "WORK EXPERIENCE", "SKILLS", "Selenium"):
        assert needle in text


def test_letter_documents():
    letter = "June 1, 2026\n\nDear Jane,\n\nPara one\ncontinues.\n\nSincerely,\nAlex"
    doc = Document(io.BytesIO(render.letter_docx(_snap(), letter)))
    paras = [p.text for p in doc.paragraphs]
    assert "Para one continues." in paras and doc.tables == []
    text = PdfReader(io.BytesIO(render.letter_pdf(_snap(), letter))).pages[0].extract_text()
    assert "Para one continues." in text


GROUPED = MASTER | {
    "headline": "Test Automation Lead",
    "skill_groups": [
        {"label": "Automation", "skills": ["Selenium", "Python"]},
        {"label": "Tools", "skills": ["Jira", "Postman", "Jenkins", "Python"]},
    ],
}


def test_reference_layout_header_and_skill_groups():
    doc = Document(
        io.BytesIO(render.resume_docx(tailor.snapshot(GROUPED, tailor.build_draft(GROUPED, []))))
    )
    paras = [p.text for p in doc.paragraphs]
    assert paras[:4] == [
        "Sam Rivera",
        "Calgary, AB, Tel.: 403-555-0100",
        "sam.jobs@example.org | https://linkedin.com/in/samrivera",
        "Test Automation Lead",
    ]
    rels = [r.target_ref for r in doc.part.rels.values() if r.reltype.endswith("/hyperlink")]
    assert rels == ["https://linkedin.com/in/samrivera"]
    assert "Automation: Selenium, Python" in paras
    assert "Tools: Python, Jira, Postman, Jenkins" in paras  # in two groups; tailored order
    assert "Other: Leadership, Agile, SQL" in paras  # skills outside every group


def test_snapshot_skill_groups_follow_hidden_and_order():
    draft = tailor.build_draft(GROUPED, [])
    for s in draft["skills"]:
        s["hidden"] = s["name"] in ("Python", "Leadership", "Agile", "SQL")
    snap = tailor.snapshot(GROUPED, draft)
    assert snap["skill_groups"] == [
        {"label": "Automation", "skills": ["Selenium"]},
        {"label": "Tools", "skills": ["Jira", "Postman", "Jenkins"]},
    ]
    # Without a headline the latest job title is used; without groups skills are one line.
    plain = tailor.snapshot(MASTER, tailor.build_draft(MASTER, []))
    assert plain["headline"] == "QA Lead"
    assert plain["skill_groups"] == [{"label": "", "skills": MASTER["skills"]}]
    # Documents generated before skill groups existed still render.
    old = {k: v for k, v in plain.items() if k not in ("skill_groups", "headline")}
    assert "Selenium" in "\n".join(
        p.text for p in Document(io.BytesIO(render.resume_docx(old))).paragraphs
    )


def test_skills_box_groups_roundtrip():
    skills, groups = model.parse_skills("Automation: Selenium, Python\nJira; SQL\nTools: Jira")
    assert skills == ["Selenium", "Python", "Jira", "SQL"]
    assert groups == [
        {"label": "Automation", "skills": ["Selenium", "Python"]},
        {"label": "Tools", "skills": ["Jira"]},
    ]
    data = model.normalise({"skills": skills, "skill_groups": groups})
    assert model.skills_text(data) == "Automation: Selenium, Python\nTools: Jira\nSQL"
    assert model.normalise({"skill_groups": groups})["skills"] == ["Selenium", "Python", "Jira"]


def test_generated_resume_imports_back():
    long = "Built a Selenium and Python framework " + "covering web, mobile and API suites " * 4
    master = GROUPED | {
        "experience": [
            GROUPED["experience"][0] | {"bullets": [long.strip(), "Mentored 5 testers."]}
        ]
        + GROUPED["experience"][1:]
    }
    snap = tailor.snapshot(master, tailor.build_draft(master, []))
    for fmt, data in (("docx", render.resume_docx(snap)), ("pdf", render.resume_pdf(snap))):
        d = importer.import_resume(data, fmt)
        assert d["headline"] == "Test Automation Lead", fmt
        first, second = d["experience"]
        assert (first["employer"], first["title"], first["location"]) == (
            "Northwind Imaging",
            "QA Lead",
            "Calgary, AB",
        ), fmt
        assert (first["start"], first["end"]) == ("Jan 2020", "Present"), fmt
        assert first["bullets"] == [long.strip(), "Mentored 5 testers."], fmt  # wrapped line joined
        assert (second["employer"], second["title"]) == (
            "Société Générale Tech",
            "Senior QA Analyst",
        )
        assert d["skill_groups"] == snap["skill_groups"], fmt
